"""讀取週報 docx：只負責取出文字，不做內容判斷。"""
from dataclasses import dataclass, field

from docx import Document

READING_LABEL = "讀經"
SERMON_LABEL = "證道"


@dataclass
class WeeklyReport:
    reading_lines: list[str] = field(default_factory=list)   # 「讀經」列第二欄，依行切分
    sermon_text: str = ""                                    # 「證道」列所有粗體 run，以空白串接


def read(path: str) -> WeeklyReport:
    report = WeeklyReport()
    sermon_runs = []
    for table in Document(path).tables:
        for row in table.rows:
            label = row.cells[0].text.strip()
            if label == READING_LABEL:
                # 有多個「讀經」列時以最後一列為準
                report.reading_lines = row.cells[1].text.strip().split("\n")
            elif label == SERMON_LABEL:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            text = run.text.strip()
                            if run.bold and text:
                                sermon_runs.append(text)
    report.sermon_text = " ".join(sermon_runs)
    return report
