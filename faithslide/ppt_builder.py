"""依 template.pptx 的範本投影片產生簡報。

範本前 5 張投影片是各類投影片的樣板，新投影片以複製樣板的方式加在後面，存檔前刪除樣板。
"""
import copy
from dataclasses import dataclass

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Length, Pt

from faithslide.bible_api import Verse
from faithslide.sermon_parser import Section

# 範本投影片索引
VERSE = 0                 # 經文
MAIN_TITLE = 1            # 主標題
AFTER_MAIN_TITLE = 2      # 接在主標題之後
MAJOR = 3                 # 大標題
OUTLINE = 4               # 中標題 / 小標題（大綱列表）
TEMPLATE_SLIDE_COUNT = 5

# 範本經文字型（微軟正黑體粗體）的字寬，以字體大小為 1；數字等寬
_GLYPH_WIDTH_EM = {
    **{digit: 0.596 for digit in "0123456789"},
    ".": 0.254, " ": 0.25, "‧": 0.254, "·": 0.255, "•": 0.439, "—": 1.08, "…": 0.859,
}
_NUMBER_GAP_EM = 0.25     # 節數與經文之間的間隔
# 開頭標點的字寬是 1 個字，但筆畫都在右半邊；懸掛到經文欄左側，讓後面的字與換行後的文字對齊
_OPENING_PUNCTUATION = "「『（《〈【〔“‘"
_HANGING_GAP_EM = 0.15    # 懸掛時節數與標點字格之間的間隔（標點字格左半邊本身是空白）

# 字體大小計算：模擬 PowerPoint 的換行，找出放得下的最大字體（以 PowerPoint 實際排版校準）
_MIN_FONT_SIZE = 20
_BASE_FONT_SIZE = 72
_FONT_SIZE_STEP = 0.5
_LINE_HEIGHT_EM = 1.2                  # 行高 = 字體大小 × 1.2
_FIT_MARGIN_EM = 0.01                  # 剛好放滿時 PowerPoint 可能因捨入而換行，保守預留
_INSET_X, _INSET_Y = Pt(7.2), Pt(3.6)  # PowerPoint 預設的文字框內距（範本未另外設定）
# 不可放在行首的結束標點：放不下時連同前一個字一起換行
_CLOSING_PUNCTUATION = set("，。、；：？！」』）》〉】〕”’…,.;:?!)")
# 不可留在行尾、會跟著下一個字換行的標點
_NO_LINE_END = set(_OPENING_PUNCTUATION)


@dataclass(frozen=True)
class TextArea:
    """經文文字框扣掉內距後的可用範圍（EMU）。"""
    width: int
    height: int


# 範本經文頁的文字框：960 × 445.75 點
DEFAULT_TEXT_AREA = TextArea(Pt(960) - 2 * _INSET_X, Pt(445.75) - 2 * _INSET_Y)


class TemplateError(Exception):
    pass


class SlideBuilder:
    def __init__(self, template_path: str):
        self.prs = Presentation(template_path)
        if len(self.prs.slides) < TEMPLATE_SLIDE_COUNT:
            raise TemplateError(f"範本投影片不足 {TEMPLATE_SLIDE_COUNT} 張：{template_path}")
        box = self.prs.slides[VERSE].shapes[1]
        self.verse_area = TextArea(box.width - 2 * _INSET_X, box.height - 2 * _INSET_Y)

    def add_main_title(self, title: str) -> None:
        slide = self._duplicate(MAIN_TITLE)
        _first_paragraph(slide.shapes[1]).runs[0].text = title
        self._duplicate(AFTER_MAIN_TITLE)

    def add_major(self, major: Section) -> None:
        slide = self._duplicate(MAJOR)
        self._set_major_heading(slide, major)

    def add_medium(self, major: Section, current: Section) -> None:
        """大綱列出 major 底下第一個到 current 為止的中標題。"""
        slide = self._duplicate(OUTLINE)
        self._set_major_heading(slide, major)
        p = _first_paragraph(slide.shapes[1])
        for i, medium in enumerate(major.children):
            p.runs[2 * i].text = f"{medium.number}."
            p.runs[2 * i + 1].text = _outline_text(medium)
            if medium is current:
                break
            p.add_run()
            p.add_run()

    def add_minor(self, major: Section, medium: Section, current: Section) -> None:
        """大綱列出所屬中標題，以及其下第一個到 current 為止的小標題。"""
        slide = self._duplicate(OUTLINE)
        self._set_major_heading(slide, major)
        p = _first_paragraph(slide.shapes[1])
        p.runs[0].text = f"{medium.number}."
        p.runs[1].text = _outline_text(medium)
        for i, minor in enumerate(medium.children, start=1):
            p.add_run()
            p.add_run()
            p.runs[2 * i].text = f"({minor.number})"
            p.runs[2 * i + 1].text = _outline_text(minor)
            if minor is current:
                break

    def add_verse(self, title: str, verse: Verse) -> bool:
        """加入經文投影片；經文太長、最小字體也放不下時回傳 False（仍會加入）。"""
        number = f"{verse.number}."
        slide = self._duplicate(VERSE)

        p = _first_paragraph(slide.shapes[0])
        if "詩篇" in title:
            title = title.replace("章", "篇")
        p.runs[0].text = title
        for run in p.runs[1:3]:
            run.text = ""

        p = _first_paragraph(slide.shapes[1])
        # 去除前後空白：例如讀經「[1]」後的 Tab 會讓第一行跳到其他定位點
        text = verse.text.replace("　", " ").strip()
        p.runs[0].text = number + "\t"
        p.runs[1].text = text
        size = calculate_font_size(number, text, self.verse_area)
        p.runs[0].font.size = size
        p.runs[1].font.size = size
        _set_hanging_indent(p, *_verse_indent(number, text, size))
        return fits(number, text, size, self.verse_area)

    def save(self, path: str) -> None:
        """刪除樣板投影片後存檔。"""
        slide_ids = self.prs.slides._sldIdLst
        for slide_id in list(slide_ids)[:TEMPLATE_SLIDE_COUNT]:
            self.prs.part.drop_rel(slide_id.rId)
            slide_ids.remove(slide_id)
        self.prs.save(path)

    def _duplicate(self, index: int):
        template = self.prs.slides[index]
        slide = self.prs.slides.add_slide(template.slide_layout)
        for shape in list(slide.shapes):
            if shape.is_placeholder:
                slide.shapes._spTree.remove(shape._element)
        for shape in template.shapes:
            slide.shapes._spTree.insert_element_before(copy.deepcopy(shape.element), "p:extLst")
        return slide

    @staticmethod
    def _set_major_heading(slide, major: Section) -> None:
        p = _first_paragraph(slide.shapes[0])
        p.runs[0].text = f"{major.number}、"
        p.runs[1].text = major.text


def _first_paragraph(shape):
    p = shape.text_frame.paragraphs[0]
    if not p.runs:
        p.add_run()
    return p


def _verse_indent(number: str, text: str, size: Length) -> tuple[int, int]:
    """回傳 (經文欄位置, Tab 位置)，單位 EMU。

    一般：節數後空一點就是經文欄，Tab 直接跳到經文欄。
    開頭是標點時（標點懸掛）：Tab 跳到節數後，標點字格放在經文欄左側，後面的字才從經文欄開始。
    """
    number_em = sum(_GLYPH_WIDTH_EM.get(char, _GLYPH_WIDTH_EM["0"]) for char in number)
    if text[:1] in _OPENING_PUNCTUATION:
        tab_em = number_em + _HANGING_GAP_EM
        column_em = tab_em + 1
    else:
        tab_em = column_em = number_em + _NUMBER_GAP_EM
    return round(size * column_em), round(size * tab_em)


def _set_hanging_indent(paragraph, column: int, tab_position: int) -> None:
    """凸排：第一行從最左邊開始放節數，Tab 跳到 tab_position；換行後的文字從 column 開始。

    範本內建的凸排是固定長度，只在原本的字體大小下對齊；字體會依經文長度調整，所以每張重新計算。
    """
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(column))
    pPr.set("indent", str(-column))
    # 範本開啟了標點懸掛：行尾的「，」會掛到右邊界外（甚至超出投影片），PowerPoint 還會壓縮整行、
    # 吃掉 Tab 的位置，讓第一行與其他行對不齊。關閉後由 wrap_lines 把標點連同前一個字移到下一行。
    pPr.set("hangingPunct", "0")
    for old in pPr.findall(qn("a:tabLst")):
        pPr.remove(old)
    tab_list = OxmlElement("a:tabLst")
    tab = OxmlElement("a:tab")
    tab.set("pos", str(tab_position))
    tab.set("algn", "l")
    tab_list.append(tab)
    pPr.insert_element_before(tab_list, "a:defRPr", "a:extLst")


def _outline_text(section: Section) -> str:
    return section.text.replace("：", "") + "\n"


def _char_width_em(char: str) -> float:
    if char in _GLYPH_WIDTH_EM:
        return _GLYPH_WIDTH_EM[char]
    if ord(char) < 0x2000:      # 英文字母等半形字元（寬度不一，取偏寬的值）
        return 0.6
    return 1.0                  # 中文字與全形標點


def wrap_lines(number: str, text: str, size: Length, area: TextArea = DEFAULT_TEXT_AREA) -> list[str]:
    """模擬 PowerPoint 的換行，回傳每一行的經文（不含節數）。

    - 第一行從 Tab 位置開始，之後各行從經文欄開始。
    - 經文段落關閉了標點懸掛：行尾放不下的結束標點（，。」等）連同前一個字一起移到下一行，
      避免下一行以標點開頭，也不會有標點跑到右邊界外。
    - 開頭標點（「等）不留在行尾，跟著下一個字換行。
    """
    column, tab = _verse_indent(number, text, size)
    limit = (area.width - tab) / size
    lines, line, used = [], "", 0.0
    for char in text:
        width = _char_width_em(char)
        if used + width <= limit - _FIT_MARGIN_EM:
            line += char
            used += width
            continue
        if char in _CLOSING_PUNCTUATION:
            # 把行尾連續的結束標點與它前面的一個字，一起移到下一行
            split = len(line)
            while split > 0 and line[split - 1] in _CLOSING_PUNCTUATION:
                split -= 1
            split = max(split - 1, 1)
        else:
            split = len(line)
            while split > 1 and line[split - 1] in _NO_LINE_END:
                split -= 1
        lines.append(line[:split])
        line = line[split:] + char
        used = sum(_char_width_em(c) for c in line)
        limit = (area.width - column) / size
    lines.append(line)
    return lines


def fits(number: str, text: str, size: Length, area: TextArea = DEFAULT_TEXT_AREA) -> bool:
    lines = len(wrap_lines(number, text, size, area))
    return lines * _LINE_HEIGHT_EM * size <= area.height


def calculate_font_size(number: str, text: str, area: TextArea = DEFAULT_TEXT_AREA) -> Pt:
    """從最大字體往下找，回傳第一個放得下的字體；最小字體仍放不下時回傳最小字體。"""
    size = _BASE_FONT_SIZE
    while size > _MIN_FONT_SIZE and not fits(number, text, Pt(size), area):
        size -= _FONT_SIZE_STEP
    return Pt(size)
