"""產生測試用週報 docx。"""
from docx import Document


def make_report(path, rows):
    """rows: [(標籤, [(文字, 是否粗體), ...] 或 純文字)]"""
    doc = Document()
    table = doc.add_table(rows=0, cols=2)
    for label, content in rows:
        cells = table.add_row().cells
        cells[0].text = label
        if isinstance(content, str):
            cells[1].text = content
        else:
            para = cells[1].paragraphs[0]
            for text, bold in content:
                para.add_run(text).bold = bold
    doc.save(path)
    return str(path)
