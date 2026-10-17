from docx import Document

from faithslide import word_reader


def make_report(path, rows):
    """rows: [(標籤, [(文字, 是否粗體), ...] 或 純文字)]"""
    doc = Document()
    table = doc.add_table(rows=0, cols=2)
    for label, content in rows:
        cells = table.add_row().cells
        cells[0].text = label
        if isinstance(content, str):
            cells[1].text = content
        else:
            para = cells[1].paragraphs[0]
            for text, bold in content:
                para.add_run(text).bold = bold
    doc.save(path)
    return str(path)


def test_read(tmp_path):
    path = make_report(tmp_path / "r.docx", [
        ("日期", "2025/10/19"),
        ("讀經", "約翰福音十五章1-2節\n[1]第一節\n[2]第二節"),
        ("證道", [("主標題 ", True), ("講員", False), (" 一、大標題 ", True), ("  ", True)]),
    ])
    report = word_reader.read(path)
    assert report.reading_lines == ["約翰福音十五章1-2節", "[1]第一節", "[2]第二節"]
    assert report.sermon_text == "主標題 一、大標題"


def test_read_missing_rows(tmp_path):
    report = word_reader.read(make_report(tmp_path / "r.docx", [("日期", "x")]))
    assert report.reading_lines == []
    assert report.sermon_text == ""
