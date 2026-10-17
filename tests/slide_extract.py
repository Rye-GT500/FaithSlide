"""把 PPT 轉成可比對的 JSON 結構（投影片版面 + 每個文字框的段落/run 文字與字體大小）。"""
from pptx import Presentation


def extract_slides(pptx_path) -> list:
    prs = Presentation(str(pptx_path))
    slides = []
    for slide in prs.slides:
        shapes = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            paragraphs = []
            for p in shape.text_frame.paragraphs:
                # 字體大小以 EMU 整數儲存，避免浮點比對誤差
                paragraphs.append([[r.text, int(r.font.size) if r.font.size else None] for r in p.runs])
            shapes.append(paragraphs)
        slides.append({"layout": slide.slide_layout.name, "shapes": shapes})
    return slides
