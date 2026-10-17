import zipfile

import pytest
from pptx.util import Pt

from faithslide import paths
from faithslide.bible_api import Verse
from faithslide.ppt_builder import SlideBuilder, TemplateError, calculate_font_size
from faithslide.sermon_parser import MAJOR, MEDIUM, MINOR, Section
from tests.slide_extract import extract_slides


def texts(slide):
    """每個文字框第一段的 run 文字。"""
    return [[run[0] for run in shape[0]] for shape in slide["shapes"]]


@pytest.fixture
def outline():
    minors = [Section(MINOR, "1", "小一：", []), Section(MINOR, "2", "小二", [])]
    mediums = [Section(MEDIUM, "1", "中一：", []), Section(MEDIUM, "2", "中二：", [], minors)]
    return Section(MAJOR, "一", "大標題：", [], mediums)


def build(tmp_path, steps):
    builder = SlideBuilder(paths.TEMPLATE_PPTX)
    steps(builder)
    out = tmp_path / "out.pptx"
    builder.save(str(out))
    return out, extract_slides(out)


def test_main_title_adds_two_slides(tmp_path):
    _, slides = build(tmp_path, lambda b: b.add_main_title("主標題"))
    assert len(slides) == 2
    assert "主標題" in texts(slides[0])[1]


def test_major(tmp_path, outline):
    _, slides = build(tmp_path, lambda b: b.add_major(outline))
    assert texts(slides[0])[0][:2] == ["一、", "大標題："]


def test_medium_lists_up_to_current(tmp_path, outline):
    _, slides = build(tmp_path, lambda b: [b.add_medium(outline, m) for m in outline.children])
    assert texts(slides[0])[0][:2] == ["一、", "大標題："]
    assert texts(slides[0])[1][:2] == ["1.", "中一\n"]
    assert texts(slides[1])[1][:4] == ["1.", "中一\n", "2.", "中二\n"]


def test_minor_lists_medium_and_minors_up_to_current(tmp_path, outline):
    medium = outline.children[1]
    _, slides = build(tmp_path, lambda b: [b.add_minor(outline, medium, m) for m in medium.children])
    assert texts(slides[0])[1][:4] == ["2.", "中二\n", "(1)", "小一\n"]
    assert texts(slides[1])[1][:6] == ["2.", "中二\n", "(1)", "小一\n", "(2)", "小二\n"]


def test_verse(tmp_path):
    _, slides = build(tmp_path, lambda b: [
        b.add_verse("約翰福音三章16節", Verse(16, "神愛世人，　甚至")),
        b.add_verse("約翰福音三章1節", Verse(1, "有一個法利賽人")),
    ])
    long_number, short_number = slides
    assert texts(long_number)[0][0] == "約翰福音三章16節"
    assert texts(long_number)[1][:2] == ["16.\t", "神愛世人， 甚至"]
    assert texts(short_number)[1][:2] == ["1.\t", "有一個法利賽人"]
    # 節數與經文使用相同字體大小（XML 以 1/100 點儲存）
    sizes = [run[1] for run in long_number["shapes"][1][0][:2]]
    assert sizes[0] == sizes[1]
    assert abs(sizes[0] - calculate_font_size("神愛世人， 甚至")) <= Pt(0.01)


def test_verse_psalm_title(tmp_path):
    _, slides = build(tmp_path, lambda b: b.add_verse("詩篇二十三章1節", Verse(1, "x")))
    assert texts(slides[0])[0][0] == "詩篇二十三篇1節"


def test_save_removes_template_slides_from_package(tmp_path):
    out, slides = build(tmp_path, lambda b: b.add_verse("約翰福音三章1節", Verse(1, "x")))
    assert len(slides) == 1
    slide_parts = [n for n in zipfile.ZipFile(out).namelist() if n.startswith("ppt/slides/slide")]
    assert len(slide_parts) == 1


def test_template_with_too_few_slides(tmp_path):
    from pptx import Presentation
    path = tmp_path / "empty.pptx"
    Presentation().save(str(path))
    with pytest.raises(TemplateError):
        SlideBuilder(str(path))


@pytest.mark.parametrize("text, expected", [
    ("", Pt(72)),
    ("短", Pt(72)),
    ("字" * 2000, Pt(20)),
])
def test_calculate_font_size_bounds(text, expected):
    assert abs(calculate_font_size(text) - expected) <= Pt(0.01)


def test_calculate_font_size_shrinks_with_length():
    sizes = [calculate_font_size("字" * n) for n in (10, 50, 100, 200)]
    assert sizes == sorted(sizes, reverse=True)
    assert sizes[0] > sizes[-1]


@pytest.mark.parametrize("number, text", [(1, "短"), (16, "字" * 40), (176, "字" * 120)])
def test_verse_hanging_indent_matches_number_width(tmp_path, number, text):
    """換行後的文字要和第一行經文對齊：凸排距離與 Tab 定位點都等於節數欄寬，且隨字體大小縮放。"""
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from faithslide.ppt_builder import _GLYPH_WIDTH_EM, _NUMBER_GAP_EM
    out, _ = build(tmp_path, lambda b: b.add_verse("標題", Verse(number, text)))
    p = Presentation(str(out)).slides[0].shapes[1].text_frame.paragraphs[0]
    pPr = p._p.find(qn("a:pPr"))
    size = calculate_font_size(text)
    expected = round(size * (sum(_GLYPH_WIDTH_EM[c] for c in f"{number}.") + _NUMBER_GAP_EM))
    assert int(pPr.get("marL")) == expected
    assert int(pPr.get("indent")) == -expected
    tabs = pPr.findall(qn("a:tabLst") + "/" + qn("a:tab"))
    assert [(int(t.get("pos")), t.get("algn")) for t in tabs] == [(expected, "l")]
